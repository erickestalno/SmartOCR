import re
from io import BytesIO
from bisect import bisect_right

import numpy as np
import pandas as pd
import streamlit as st
import importlib
from PIL import Image, ImageOps
from docx import Document
from fpdf import FPDF

# ==============================================================================
# ImageScan - Smart OCR Studio
# 100% offline (RapidOCR / ONNX): tanpa API key, tanpa internet, ringan & cepat.
# ==============================================================================


# ------------------------------------------------------------------------------
# MESIN OCR
# ------------------------------------------------------------------------------
@st.cache_resource(show_spinner="Menyiapkan mesin OCR (hanya sekali)...")
def load_ocr():
    galat = []
    try:
        from rapidocr import RapidOCR  # paket baru (Python 3.8 - 3.13+)
        return "baru", RapidOCR()
    except Exception as e:
        galat.append(f"rapidocr -> {type(e).__name__}: {e}")
    try:
        RapidOCR = importlib.import_module("rapidocr_onnxruntime").RapidOCR  # paket lama (Python <= 3.12)
        return "lama", RapidOCR()
    except Exception as e:
        galat.append(f"rapidocr_onnxruntime -> {type(e).__name__}: {e}")
    raise RuntimeError("\n".join(galat))


def siapkan_gambar(img, sisi_maks=2200):
    """Putar sesuai EXIF (foto HP), jadikan RGB, kecilkan bila terlalu besar (biar cepat)."""
    img = ImageOps.exif_transpose(img).convert("RGB")
    if max(img.size) > sisi_maks:
        img.thumbnail((sisi_maks, sisi_maks), Image.LANCZOS)
    return img


def jalankan_ocr(img):
    """Mengembalikan list (kotak 4 titik, teks, skor)."""
    jenis, mesin = load_ocr()
    arr = np.array(img)
    if jenis == "baru":
        out = mesin(arr)
        if out is None or out.boxes is None:
            return []
        return [(b, t, float(s)) for b, t, s in zip(out.boxes, out.txts, out.scores)]
    hasil, _ = mesin(arr)
    return [(b, t, float(s)) for b, t, s in (hasil or [])]


# ------------------------------------------------------------------------------
# MENYUSUN TEKS: BARIS & KOLOM BERDASARKAN POSISI
# ------------------------------------------------------------------------------
_SIMBOL = re.compile(r"^[\|\[\]\(\)\{\}_\-—–=~\.\s]+$")  # sisa garis tabel yang terbaca sebagai teks


def _kumpulkan_item(hasil_ocr):
    items = []
    for bbox, text, skor in hasil_ocr:
        text = str(text).strip()
        if not text or skor < 0.3 or _SIMBOL.match(text):
            continue
        xs = [float(p[0]) for p in bbox]
        ys = [float(p[1]) for p in bbox]
        w, h = max(xs) - min(xs), max(ys) - min(ys)
        if len(text) <= 2 and set(text) <= set("|Il!") and h > 2.5 * max(w, 1):
            continue  # garis tegak yang terbaca sebagai huruf
        items.append({"x0": min(xs), "x1": max(xs), "yc": (min(ys) + max(ys)) / 2, "h": h, "text": text})
    return items


def _kelompok_baris(items, med_h):
    items = sorted(items, key=lambda i: i["yc"])
    baris = []
    for it in items:
        if baris and abs(it["yc"] - baris[-1]["yc"]) <= med_h * 0.6:
            baris[-1]["items"].append(it)
            baris[-1]["yc"] = float(np.mean([i["yc"] for i in baris[-1]["items"]]))
        else:
            baris.append({"yc": it["yc"], "items": [it]})
    return [b["items"] for b in baris]


def _cari_pemisah_kolom(items, med_h, n_baris):
    """
    Pemisah kolom = celah vertikal yang hampir tidak dilewati teks di semua baris.
    (Cara ini tahan terhadap garis tabel tebal & jarak antar kata yang sempit.)
    """
    if n_baris < 4:
        return []
    x_min = int(min(i["x0"] for i in items))
    x_max = int(max(i["x1"] for i in items))
    cov = np.zeros(x_max - x_min + 2, dtype=int)
    pad = med_h * 0.2  # kotak deteksi biasanya sedikit lebih lebar dari hurufnya
    for it in items:
        a, b = it["x0"] + pad, it["x1"] - pad
        if b < a:
            a = b = (it["x0"] + it["x1"]) / 2
        cov[int(a) - x_min: int(b) - x_min + 1] += 1

    toleransi = 0 if n_baris < 8 else max(1, int(n_baris * 0.1))  # baris judul yang melintang
    lebar_min = max(2, int(med_h * 0.25))
    kosong = cov <= toleransi

    pemisah, i, n = [], 0, len(kosong)
    while i < n:
        if kosong[i]:
            j = i
            while j < n and kosong[j]:
                j += 1
            if i > 0 and j < n and (j - i) >= lebar_min:
                pemisah.append(x_min + (i + j) / 2)
            i = j
        else:
            i += 1
    return pemisah


def _pecah_teks(text, x0, x1, pemisah_dalam):
    """Bila satu kotak terlanjur menggabung 2 kolom, pecah di spasi terdekat dari posisi celah."""
    potongan, sisa, kiri = [], text, x0
    for g in pemisah_dalam:
        spasi = [m.start() for m in re.finditer(r"\s", sisa)]
        if not spasi or x1 - kiri <= 0:
            return None
        idx = (g - kiri) / (x1 - kiri) * len(sisa)
        dekat = min(spasi, key=lambda k: abs(k - idx))
        if abs(dekat - idx) > max(2, 0.25 * len(sisa)):
            return None
        potongan.append(sisa[:dekat].strip())
        sisa, kiri = sisa[dekat:].strip(), g
    potongan.append(sisa)
    return potongan if all(potongan) else None


def _baris_judul(row, med_h):
    """Baris judul/kategori (mis. 'NAL & TOOTH CARE 8') yang terpotong garis tabel -> jadikan satu sel."""
    its = sorted(row, key=lambda i: i["x0"])
    if len(its) < 2:
        return False
    if any(b["x0"] - a["x1"] >= med_h * 0.8 for a, b in zip(its, its[1:])):
        return False
    return not any(sum(c.isdigit() for c in i["text"]) / max(len(i["text"]), 1) > 0.6 for i in its)


def _taruh(cells, idx, text):
    cells[idx] = (cells[idx] + " " + text).strip()


def susun_baris(hasil_ocr, mode_tabel=True):
    """Hasil OCR -> list baris; tiap baris = list sel (kolom). Sel kosong dipertahankan agar kolom sejajar."""
    items = _kumpulkan_item(hasil_ocr)
    if not items:
        return []
    med_h = float(np.median([i["h"] for i in items])) or 1.0
    baris = _kelompok_baris(items, med_h)
    pemisah = _cari_pemisah_kolom(items, med_h, len(baris)) if mode_tabel else []

    hasil = []
    for row in baris:
        cells = [""] * (len(pemisah) + 1)
        if pemisah and _baris_judul(row, med_h):
            urut = sorted(row, key=lambda i: i["x0"])
            _taruh(cells, bisect_right(pemisah, urut[0]["x0"] + med_h * 0.2), " ".join(i["text"] for i in urut))
            hasil.append(cells)
            continue
        for it in sorted(row, key=lambda i: i["x0"]):
            kiri = it["x0"] + med_h * 0.2  # abaikan lebihan padding kotak deteksi
            kol = bisect_right(pemisah, kiri)
            dalam = [g for g in pemisah if kiri + 2 < g < it["x1"] - 2]
            ada_kiri = any(o is not it and o["x0"] < it["x0"] for o in row)
            if dalam and ada_kiri:
                potong = _pecah_teks(it["text"], it["x0"], it["x1"], dalam)
                if potong and len(potong) == len(dalam) + 1:
                    for k, p in enumerate(potong):
                        _taruh(cells, kol + k, p)
                    continue
            _taruh(cells, kol, it["text"])
        hasil.append(cells)
    return hasil


def baris_ke_teks(rows):
    return "\n".join(" | ".join(r) for r in rows)


def teks_ke_baris(text):
    """Satu baris teks = satu baris tabel; kolom dipisah tanda '|'."""
    rows = []
    for line in text.split("\n"):
        line = line.strip()
        if not line or re.fullmatch(r"[\|\-:\s+=]+", line):
            continue
        if "|" in line:
            if line.startswith("|") and line.endswith("|"):  # gaya markdown
                line = line[1:-1]
            rows.append([c.strip() for c in line.split("|")])
        else:
            rows.append([line])
    return rows


def buat_df(rows, num_cols, headers):
    tetap = []
    for r in rows:
        if len(r) > num_cols:  # kelebihan sel digabung ke kolom terakhir
            r = r[: num_cols - 1] + [" ".join(x for x in r[num_cols - 1:] if x)]
        tetap.append(r + [""] * (num_cols - len(r)))
    return pd.DataFrame(tetap, columns=headers)


# ------------------------------------------------------------------------------
# EXPORT
# ------------------------------------------------------------------------------
def buat_pdf(text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    aman = text.encode("latin-1", "replace").decode("latin-1")
    pdf.multi_cell(0, 8, aman, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def buat_docx(text):
    doc = Document()
    doc.add_heading("Hasil Ekstraksi Teks", 0)
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def buat_xlsx(df):
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data_Scan")
    return buf.getvalue()


def tampil_gambar(img):
    try:
        st.image(img, width="stretch")
    except Exception:  # Streamlit versi lama
        st.image(img, use_container_width=True)


def tampil_tabel(df):
    try:
        return st.data_editor(df, num_rows="dynamic", width="stretch")
    except Exception:
        return st.data_editor(df, num_rows="dynamic", use_container_width=True)


# ===== ANTARMUKA =====
st.set_page_config(page_title="Scan - Smart OCR Studio", layout="wide")

st.markdown("""
<style>
    .stApp { background: linear-gradient(135deg, #f5f7fa 0%, #c3cfe2 100%); }
    .stButton>button {
        background: linear-gradient(90deg, #7F00FF 0%, #E100FF 100%);
        color: white; border: none; border-radius: 12px;
        font-weight: bold; padding: 10px 24px; transition: all 0.3s ease;
    }
    .stButton>button:hover { transform: scale(1.03); }
</style>
""", unsafe_allow_html=True)

st.title("‎𝄃𝄂𝄂𝄀𝄁𝄃𝄂𝄂𝄃𝄂𝄃𝄂𝄃𝄂𝄂 ImageScan - Smart OCR Studio")
st.caption("Pindai teks dari gambar, susun baris & kolom secara rapi, lalu ekspor ke Excel/Word/PDF - tanpa internet.")

try:
    load_ocr()  # dimuat sekali di awal, jadi tombol Pindai langsung cepat
except Exception as e:
    st.error("Mesin OCR gagal dimuat. Salin pesan di bawah ini untuk dicek:")
    st.code(str(e))
    st.stop()

# Sidebar
st.sidebar.header("⚙️ Pengaturan Studio")
mode_input = st.sidebar.radio("Sumber Gambar:", ["Upload File", "Kamera HP / Webcam"])
mode_crop = st.sidebar.radio("Mode Ekstraksi:", ["Seluruh Teks", "Pilih Area Teks (Crop)"])
mode_susun = st.sidebar.radio(
    "Susunan Hasil:",
    ["Tabel / daftar berkolom", "Teks biasa (per baris)"],
    help="Pilih 'Tabel' untuk daftar barang/nota/tabel. Pilih 'Teks biasa' untuk paragraf.",
)

if "teks" not in st.session_state:
    st.session_state.teks = ""
if "scan_id" not in st.session_state:
    st.session_state.scan_id = 0

# 1. Input gambar
image = None
if mode_input == "Upload File":
    f = st.file_uploader("Unggah foto atau screenshot dokumen kamu", type=["jpg", "jpeg", "png"])
    if f:
        image = siapkan_gambar(Image.open(f))
else:
    f = st.camera_input("Ambil foto dokumen secara langsung")
    if f:
        image = siapkan_gambar(Image.open(f))

if image:
    col_img, col_out = st.columns([1, 1])

    with col_img:
        st.subheader("🖼️ Gambar Sumber")
        if mode_crop == "Pilih Area Teks (Crop)":
            try:
                from streamlit_cropper import st_cropper
                st.info("Tarik kotak ungu untuk memilih area teks yang ingin dipindai:")
                target_img = st_cropper(image, realtime_update=True, box_color="#7F00FF", aspect_ratio=None)
            except Exception:
                st.info("Geser slider untuk menentukan area teks yang ingin dipindai:")
                w, h = image.size
                left = st.slider("Posisi Kiri", 0, max(1, w - 10), 0)
                top = st.slider("Posisi Atas", 0, max(1, h - 10), 0)
                right = st.slider("Posisi Kanan", left + 10, w, w)
                bottom = st.slider("Posisi Bawah", top + 10, h, h)
                target_img = image.crop((left, top, right, bottom))
                tampil_gambar(target_img)
        else:
            tampil_gambar(image)
            target_img = image

    # 2. Pindai
    if st.button("🔍 Pindai Teks Sekarang"):
        with st.spinner("Membaca teks dari gambar..."):
            try:
                hasil = jalankan_ocr(target_img)
                rows = susun_baris(hasil, mode_tabel=(mode_susun == "Tabel / daftar berkolom"))
                st.session_state.teks = baris_ke_teks(rows)
                st.session_state.scan_id += 1
                if not st.session_state.teks.strip():
                    st.warning("Tidak ada teks yang terbaca. Coba gambar yang lebih jelas/terang atau crop area teksnya.")
            except Exception as e:
                st.error(f"Gagal memindai: {e}")

    # 3. Pratinjau & edit
    with col_out:
        st.subheader("📝 Pratinjau & Edit Teks")
        st.caption("Satu baris = satu baris tabel. Tanda ' | ' memisahkan kolom.")
        edited_text = st.text_area("Edit teks hasil scan di sini:", key="teks", height=300)

    st.markdown("---")

    # 4. Tabel Excel
    st.subheader("📊 Susun Data ke Tabel Excel")
    rows = teks_ke_baris(edited_text)
    pakai_header = st.checkbox("Pakai baris pertama sebagai header kolom", value=False)
    data_rows = rows[1:] if (pakai_header and len(rows) > 1) else rows
    auto_cols = min(10, max((len(r) for r in rows), default=1))
    sid = f"{st.session_state.scan_id}_{int(pakai_header)}"

    num_cols = st.number_input("Jumlah Header Kolom:", min_value=1, max_value=10,
                               value=auto_cols, key=f"ncols_{sid}")
    col_names = []
    for i, c in enumerate(st.columns(num_cols)):
        awal = f"Kolom_{i + 1}"
        if pakai_header and rows and i < len(rows[0]) and rows[0][i]:
            awal = rows[0][i]
        col_names.append(c.text_input(f"Header {i + 1}:", value=awal, key=f"hdr_{sid}_{i}"))

    df = buat_df(data_rows, num_cols, col_names)
    st.write("Kelola dan lengkapi isi tabel langsung di bawah ini:")
    edited_df = tampil_tabel(df)

    # 5. Export
    st.markdown("---")
    st.subheader("💾 Export Hasil")
    c1, c2, c3, c4 = st.columns(4)

    c1.download_button("📥 Download TXT", data=edited_text, file_name="hasil_scan.txt", mime="text/plain")
    c2.download_button("📥 Download Excel (.xlsx)", data=buat_xlsx(edited_df), file_name="hasil_scan.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    c3.download_button("📥 Download Word (.docx)", data=buat_docx(edited_text), file_name="hasil_scan.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    try:
        c4.download_button("📥 Download PDF (.pdf)", data=buat_pdf(edited_text), file_name="hasil_scan.pdf",
                           mime="application/pdf")
    except Exception as e:
        c4.warning(f"PDF gagal dibuat: {e}")
